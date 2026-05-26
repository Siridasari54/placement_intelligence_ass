import os
from typing import Dict, Any
import docx
from app.utils.logger import logger

class DocxLoader:
    def __init__(self, file_path: str):
        self.file_path = file_path
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

    def load(self) -> Dict[str, Any]:
        """Loads a .docx file and returns text and metadata."""
        try:
            doc = docx.Document(self.file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
                    
            # Extract tables as well
            tables_data = []
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(row_cells)
                tables_data.append(table_rows)
                
            text_content = "\n".join(full_text)
            logger.info(f"Successfully loaded DOCX: {self.file_path}")
            return {
                "text": text_content,
                "tables": tables_data,
                "metadata": {
                    "source": self.file_path,
                    "type": "docx"
                }
            }
        except Exception as e:
            logger.error(f"Error loading DOCX {self.file_path}: {e}")
            raise e
